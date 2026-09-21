"""
SUMMARY 차트(엑셀)를 새로 만들거나 기존 SUMMARY 파일에 새 제출 데이터를 병합해서
다시 만들어주는 모듈.

컬럼 구성 (스타일/컬러웨이별로 가장 최근 제출분 한 줄만 표시, 내부/오픈을 각각 표시):
  Season | Style/Colour | Style Description | Stage |
  내부 FOB | 내부 CM | 내부 Margin % | 오픈 FOB | 오픈 CM | REMARK

CM은 FOB 안에 포함되어 있는 원가 구성요소이므로(FOB = 원부자재비 등 기타 고정비용 + CM),
CM을 수정하면 그만큼 FOB도 같이 움직인다. 내부FOB는 내부CM에, 오픈FOB는 오픈CM에만
연동된다 (내부CM을 바꾸면 내부FOB만, 오픈CM을 바꾸면 오픈FOB만 바뀐다). 이를 위해
'계산정보' 시트에 스타일별 기준값(생성 시점의 내부/오픈 FOB·CM)을 저장해두고,
Summary의 FOB 칸은 그 기준값과 현재 CM 칸을 참조하는 엑셀 수식으로 만든다:
  FOB_new = FOB_기준값 + (CM_new - CM_기준값)
내부MARGIN%는 원본 파일의 실제 수식과 동일하게 '1 - (내부FOB / 오픈FOB)'로 계산한다.
그래서 내부CM뿐 아니라 오픈CM을 바꿔도(분모인 오픈FOB가 바뀌므로) 내부MARGIN%가
같이 재계산된다. 오픈 MARGIN%는 원본 파일에 별도 셀 자체가 없어(=계산 로직을 알 수
없어) 컬럼을 두지 않는다.
"""
import io
import re
from dataclasses import dataclass
from typing import Optional

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = "FF0070C0"
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
EDITABLE_FILL = PatternFill("solid", fgColor="FFFFF2CC")
FONT_NAME = "Calibri"

CALC_SHEET_NAME = "계산정보"


@dataclass
class SummaryRow:
    season: str
    style_no: str
    colorway: str  # "SOLID" / "HEATHER" / "PRINT" / "" (컬러웨이 구분 없는 단일 항목)
    description: str
    stage: str = ""  # 참고용: 이 값이 어느 단계(P1/PP/SMS 등)에서 나온 것인지
    fob_internal: Optional[float] = None
    cm_internal: Optional[float] = None
    margin_internal: Optional[float] = None
    fob_open: Optional[float] = None
    cm_open: Optional[float] = None
    remark: str = ""

    @property
    def key(self):
        return (self.style_no, self.colorway)

    @property
    def style_cell_text(self):
        return f"{self.style_no} / {self.colorway}" if self.colorway else self.style_no


def _norm_header(text) -> str:
    return re.sub(r"\s+", "", str(text).strip().upper()) if text is not None else ""


def _parse_style_cell(text: str):
    if text is None:
        return "", ""
    text = str(text).strip()
    if "/" in text:
        style_no, colorway = text.split("/", 1)
        return style_no.strip(), colorway.strip()
    return text, ""


_GENDER_RE = re.compile(r"^(M|W)\b\s*(.*)$")


def _natural_key(text: str) -> str:
    """숫자가 포함된 텍스트를 자연스러운 순서(3.5 < 5 < 7 등)로 정렬하기 위한 키.
    숫자 부분을 고정 폭으로 0-패딩해서 문자열 비교만으로 숫자 순서가 맞도록 한다."""
    def pad(m):
        try:
            return f"{float(m.group(0)):020.4f}"
        except ValueError:
            return m.group(0)
    return re.sub(r"\d+(?:\.\d+)?", pad, (text or "").strip().upper())


def _desc_sort_key(row: "SummaryRow"):
    """Style Description을 비슷한 이름끼리(예: Race Day류, Glide Short 3.5/5/7류)
    묶어서 정렬하기 위한 키. 성별(M/W) 접두어는 떼어내 그룹핑 기준에서 빼고,
    같은 이름 그룹 안에서는 M을 W보다 위에 오도록 한다."""
    desc = (row.description or "").strip()
    m = _GENDER_RE.match(desc)
    if m:
        gender, rest = m.group(1), m.group(2)
    else:
        gender, rest = "", desc
    gender_rank = {"M": 0, "W": 1}.get(gender, 2)
    return (_natural_key(rest), gender_rank, row.style_no)


def sort_rows(rows: list) -> list:
    """Style Description 기준으로 비슷한 이름끼리 묶고, 같은 그룹 안에서는
    M(남성) 항목을 W(여성) 항목보다 위로 오도록 정렬한다."""
    return sorted(rows, key=_desc_sort_key)


def parse_existing_summary(file_like) -> list[SummaryRow]:
    """기존 SUMMARY 엑셀에서 행 데이터를 읽어 SummaryRow 리스트로 반환.

    내부/오픈으로 나뉜 새 방식 헤더("내부FOB", "오픈CM" 등)와, 그 이전의 단일
    FOB/CM/MARGIN% 방식, 그리고 더 예전의 날짜/단계별 여러 컬럼 방식을 모두
    최대한 인식한다. 예전 방식(그룹이 여러 개)이면 값이 채워진 마지막(=가장 최근)
    그룹의 값을 대표값으로 사용하고, 내부/오픈 구분이 없던 파일은 두 쪽 다 같은
    값으로 채운다."""
    wb = openpyxl.load_workbook(file_like, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = None
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, 6):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().upper() == "SEASON":
                header_row = r
                break
        if header_row:
            break
    if header_row is None:
        return []

    col_headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str):
            col_headers[c] = _norm_header(v)

    season_col = next((c for c, v in col_headers.items() if v == "SEASON"), 1)
    style_col = next((c for c, v in col_headers.items() if v in ("STYLE", "STYLE/COLOUR", "STYLE/COLOR")), None)
    desc_col = next((c for c, v in col_headers.items() if v in ("DESCRIPTION", "STYLEDESCRIPTION")), None)
    stage_col = next((c for c, v in col_headers.items() if v == "STAGE"), None)
    remark_col = next((c for c, v in col_headers.items() if v == "REMARK"), None)

    if style_col is None or desc_col is None:
        return []

    internal_fob_col = next((c for c, v in col_headers.items() if v in ("내부FOB",)), None)
    internal_cm_col = next((c for c, v in col_headers.items() if v in ("내부CM",)), None)
    internal_margin_col = next((c for c, v in col_headers.items() if v in ("내부MARGIN%",)), None)
    open_fob_col = next((c for c, v in col_headers.items() if v in ("오픈FOB",)), None)
    open_cm_col = next((c for c, v in col_headers.items() if v in ("오픈CM",)), None)

    # 내부/오픈 구분이 없던 예전 파일: 단일 FOB/CM/MARGIN% 컬럼(들)을 찾아 양쪽에 같이 채운다.
    legacy_fob_cols = sorted(c for c, v in col_headers.items() if v.startswith("FOB") and v != "내부FOB" and v != "오픈FOB")

    def legacy_group(fob_col):
        cm_col = fob_col + 1 if col_headers.get(fob_col + 1) == "CM" else None
        margin_col = fob_col + 2 if cm_col and "MARGIN" in col_headers.get(fob_col + 2, "") else None
        return {"fob_col": fob_col, "cm_col": cm_col, "margin_col": margin_col}

    legacy_groups = [legacy_group(c) for c in legacy_fob_cols]

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        style_cell = ws.cell(row=r, column=style_col).value
        description = ws.cell(row=r, column=desc_col).value
        if not style_cell and not description:
            continue
        style_no, colorway = _parse_style_cell(style_cell)
        season = ws.cell(row=r, column=season_col).value
        stage_val = ws.cell(row=r, column=stage_col).value if stage_col else ""
        remark_val = ws.cell(row=r, column=remark_col).value if remark_col else ""
        stage_label = str(stage_val).strip() if stage_val else ""

        def numval(col):
            if not col:
                return None
            v = ws.cell(row=r, column=col).value
            return v if isinstance(v, (int, float)) else None

        if internal_fob_col or internal_cm_col or internal_margin_col or open_fob_col or open_cm_col:
            fob_internal = numval(internal_fob_col)
            cm_internal = numval(internal_cm_col)
            margin_internal = numval(internal_margin_col)
            fob_open = numval(open_fob_col) if open_fob_col else fob_internal
            cm_open = numval(open_cm_col) if open_cm_col else cm_internal
        else:
            fob_internal = cm_internal = margin_internal = None
            for g in legacy_groups:
                gf, gc, gm = numval(g["fob_col"]), numval(g["cm_col"]), numval(g["margin_col"])
                if gf is None and gc is None:
                    continue
                if gf is not None:
                    fob_internal = gf
                if gc is not None:
                    cm_internal = gc
                if gm is not None:
                    margin_internal = gm
            fob_open, cm_open = fob_internal, cm_internal

        rows.append(
            SummaryRow(
                season=season or "",
                style_no=style_no,
                colorway=colorway,
                description=description or "",
                stage=stage_label,
                fob_internal=fob_internal,
                cm_internal=cm_internal,
                margin_internal=margin_internal,
                fob_open=fob_open,
                cm_open=cm_open,
                remark=remark_val or "",
            )
        )
    return rows


def merge_rows(existing: list[SummaryRow], new: list[SummaryRow]) -> list[SummaryRow]:
    """기존 행과 새로 파싱한 행을 style_no+colorway 기준으로 병합.
    같은 스타일/컬러웨이가 이미 있으면 새로 추출한(더 최근 제출분) 값으로 덮어쓴다."""
    by_key = {r.key: r for r in existing}
    ordered_keys = [r.key for r in existing]
    for nr in new:
        if nr.key in by_key:
            er = by_key[nr.key]
            if nr.fob_internal is not None:
                er.fob_internal = nr.fob_internal
            if nr.cm_internal is not None:
                er.cm_internal = nr.cm_internal
            if nr.margin_internal is not None:
                er.margin_internal = nr.margin_internal
            if nr.fob_open is not None:
                er.fob_open = nr.fob_open
            if nr.cm_open is not None:
                er.cm_open = nr.cm_open
            if nr.stage:
                er.stage = nr.stage
            if nr.description:
                er.description = nr.description
            if nr.season:
                er.season = nr.season
            if nr.remark:
                er.remark = nr.remark
        else:
            by_key[nr.key] = nr
            ordered_keys.append(nr.key)
    return [by_key[k] for k in ordered_keys]


calc_ws_note = (
    "※ 이 시트는 계산용입니다. 직접 값을 고치지 마세요. Summary 시트 생성 시점의 "
    "FOB·CM 기준값을 저장해두고, Summary의 FOB 칸이 이 기준값과 현재 CM 칸을 참조하는 "
    "수식으로 계산되게 합니다 (CM은 FOB 안에 포함된 항목이라, CM이 바뀌면 그만큼 FOB도 "
    "같이 움직입니다. 내부CM은 내부FOB만, 오픈CM은 오픈FOB만 움직입니다). 내부MARGIN%는 "
    "'1-(내부FOB/오픈FOB)'로 계산되므로, 내부CM과 오픈CM 중 어느 쪽을 바꾸어도 함께 재계산됩니다."
)


def build_summary_workbook(rows: list[SummaryRow], season_label: str, brand: str):
    rows = sort_rows(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"

    headers = [
        "Season", "Style/Colour", "Style Description", "Stage",
        "내부FOB", "내부CM", "내부MARGIN%",
        "오픈FOB", "오픈CM",
        "REMARK",
    ]
    header_row = 1
    data_start_row = 2

    def style_header(cell):
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        cell.fill = HEADER_FILL

    for i, name in enumerate(headers):
        c = ws.cell(row=header_row, column=i + 1, value=name)
        style_header(c)

    calc_ws = wb.create_sheet(CALC_SHEET_NAME)
    CALC_HEADERS = [
        "Style/Colour", "내부FOB(기준값)", "내부CM(기준값)",
        "오픈FOB(기준값)", "오픈CM(기준값)",
    ]
    for i, name in enumerate(CALC_HEADERS):
        c = calc_ws.cell(row=header_row, column=i + 1, value=name)
        c.font = Font(name=FONT_NAME, bold=True, size=9)
    calc_ws.cell(row=header_row, column=len(CALC_HEADERS) + 2, value=calc_ws_note)
    calc_ws.cell(row=header_row, column=len(CALC_HEADERS) + 2).font = Font(name=FONT_NAME, size=9, italic=True, color="FF808080")
    for col, w in {1: 20, 2: 14, 3: 14, 4: 14, 5: 14, 7: 90}.items():
        calc_ws.column_dimensions[get_column_letter(col)].width = w
    CALC_COL_STYLE, CALC_COL_FOB_I, CALC_COL_CM_I, CALC_COL_FOB_O, CALC_COL_CM_O = 1, 2, 3, 4, 5

    currency_fmt = '_(\\$* #,##0.00_);_(\\$* \\(#,##0.00\\);_(\\$* "-"??_);_(@_)'

    r = data_start_row
    review_flags = []  # (row, col, message) - UI에 검토 필요 항목 표시용
    for row_data in rows:
        c_season = ws.cell(row=r, column=1, value=row_data.season)
        c_style = ws.cell(row=r, column=2, value=row_data.style_cell_text)
        c_desc = ws.cell(row=r, column=3, value=row_data.description)
        c_stage = ws.cell(row=r, column=4, value=row_data.stage)
        for cell in (c_season, c_style, c_desc, c_stage):
            cell.font = Font(name=FONT_NAME, size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

        fob_i_cell = ws.cell(row=r, column=5)
        fob_i_cell.number_format = currency_fmt
        fob_i_cell.font = Font(name=FONT_NAME, size=9, bold=True, color=BLUE)
        fob_i_cell.alignment = Alignment(horizontal="center", vertical="center")
        fob_i_cell.border = BORDER

        # 내부 CM (F열) - 사용자가 직접 수정하는 칸.
        cm_i_cell = ws.cell(row=r, column=6, value=row_data.cm_internal)
        cm_i_cell.font = Font(name=FONT_NAME, size=9, bold=True)
        cm_i_cell.alignment = Alignment(horizontal="center", vertical="center")
        cm_i_cell.border = BORDER
        cm_i_cell.fill = EDITABLE_FILL

        fob_o_cell = ws.cell(row=r, column=8)
        fob_o_cell.number_format = currency_fmt
        fob_o_cell.font = Font(name=FONT_NAME, size=9, bold=True, color=BLUE)
        fob_o_cell.alignment = Alignment(horizontal="center", vertical="center")
        fob_o_cell.border = BORDER

        margin_i_cell = ws.cell(row=r, column=7)
        margin_i_cell.number_format = "0.00%"
        margin_i_cell.font = Font(name="맑은 고딕", size=9)
        margin_i_cell.alignment = Alignment(horizontal="center", vertical="center")
        margin_i_cell.border = BORDER

        # 오픈 CM (I열) - 사용자가 직접 수정하는 칸.
        cm_o_cell = ws.cell(row=r, column=9, value=row_data.cm_open)
        cm_o_cell.font = Font(name=FONT_NAME, size=9)
        cm_o_cell.alignment = Alignment(horizontal="center", vertical="center")
        cm_o_cell.border = BORDER

        # 계산정보 시트에 이 행의 생성 시점 기준값(FOB·CM)을 저장해두고, Summary의
        # FOB/MARGIN% 칸은 그 기준값과 현재 CM 칸을 참조하는 수식으로 만든다.
        # CM은 FOB 안에 포함된 원가 항목이므로, CM이 바뀌면 그 변동분만큼 FOB도 같이
        # 움직인다: FOB_new = FOB_기준값 + (CM_new - CM_기준값). 내부/오픈은 서로
        # 독립적으로 움직인다 (내부CM은 내부FOB만, 오픈CM은 오픈FOB만 움직인다).
        calc_ws.cell(row=r, column=CALC_COL_STYLE, value=row_data.style_cell_text)

        has_internal_baseline = row_data.fob_internal is not None and row_data.cm_internal is not None
        if has_internal_baseline:
            calc_ws.cell(row=r, column=CALC_COL_FOB_I, value=row_data.fob_internal)
            calc_ws.cell(row=r, column=CALC_COL_CM_I, value=row_data.cm_internal)
            fob_i_cell.value = (
                f"='{CALC_SHEET_NAME}'!{get_column_letter(CALC_COL_FOB_I)}{r}"
                f"+(F{r}-'{CALC_SHEET_NAME}'!{get_column_letter(CALC_COL_CM_I)}{r})"
            )
        else:
            fob_i_cell.value = row_data.fob_internal

        has_open_baseline = row_data.fob_open is not None and row_data.cm_open is not None
        if has_open_baseline:
            calc_ws.cell(row=r, column=CALC_COL_FOB_O, value=row_data.fob_open)
            calc_ws.cell(row=r, column=CALC_COL_CM_O, value=row_data.cm_open)
            fob_o_cell.value = (
                f"='{CALC_SHEET_NAME}'!{get_column_letter(CALC_COL_FOB_O)}{r}"
                f"+(I{r}-'{CALC_SHEET_NAME}'!{get_column_letter(CALC_COL_CM_O)}{r})"
            )
            cm_o_cell.fill = EDITABLE_FILL
        else:
            fob_o_cell.value = row_data.fob_open

        # 내부MARGIN% = 1 - (내부FOB / 오픈FOB). 원본 파일의 실제 마진% 수식과
        # 동일한 구조라, 내부CM뿐 아니라 오픈CM을 바꿔도(분모인 오픈FOB가 바뀌므로)
        # 함께 재계산된다. 이 계산에는 내부FOB·오픈FOB 두 칸이 모두 값을 갖고
        # 있어야 하므로(수식이든 고정값이든), 둘 다 있을 때만 수식으로 넣는다.
        if row_data.fob_internal is not None and row_data.fob_open is not None:
            margin_i_cell.value = f"=1-(E{r}/H{r})"
        else:
            margin_i_cell.value = row_data.margin_internal
            if row_data.fob_internal is not None:
                review_flags.append((r, 7, f"{row_data.style_cell_text}: 내부MARGIN% 자동 재계산 정보 부족 (수동 확인 필요)"))

        remark_cell = ws.cell(row=r, column=10, value=row_data.remark)
        remark_cell.font = Font(name=FONT_NAME, size=9)
        remark_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        remark_cell.border = BORDER

        r += 1

    widths = {1: 9, 2: 18, 3: 26, 4: 9, 5: 11, 6: 9, 7: 11, 8: 11, 9: 9, 10: 42}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = ws.cell(row=data_start_row, column=5)

    return wb, review_flags


def workbook_to_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
