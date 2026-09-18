"""
SUBMITTED 원가(costing) 엑셀 파일을 읽어서 SUMMARY 차트에 필요한 정보를 추출하는 모듈.

추출 대상 (사용자 지침 기준):
  2. SUBMITTED로 끝나는 파일에서 정보 추출
  3. 파일명 시작 부분: 시즌(S27/F27/S28...) + 7자리 스타일넘버 + 영문 DESCRIPTION
  4. 제출일자/가격은 각 시트의 "CBS + 날짜 + 가격"에서, CM은 좌/우 2개 중 좌측 값 사용
  5. 시트가 2개 초과면 날짜별 비교(표+그래프)
  6. 우측 REMARK(노란 셀) 내용을 요약의 REMARK 란에 반영 + 단계/날짜별 차이점도 함께 기재
"""
import re
import datetime as dt
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

YELLOW_FILLS = {"FFFFFF00", "00FFFF00", "FFFFFF99"}  # 실 사용 노란색 계열

STOPWORDS = re.compile(
    r"^(fab ?trim|fabtrim|fab|trim|updated|update|submitted|submit|v|ver\d*|"
    r"draft|final|rev\d*|revised|p\d+)$",
    re.IGNORECASE,
)

# 시트 탭 이름에 적혀있는 샘플 단계 키워드. 왼쪽에 나올수록(문자열 상 더 앞쪽 위치)
# 우선한다 (예: 'GTM2-F27 PP 6.03' -> GTM2가 PP보다 앞에 있으므로 GTM2로 판단).
STAGE_KEYWORDS = ["GTM2", "GTM1", "SMS", "PP", "P1", "P2", "P3"]
STAGE_ORDER = {name: i for i, name in enumerate(["P1", "P2", "P3", "GTM1", "GTM2", "PP", "SMS"])}


SEASON_TOKEN_RE = re.compile(r"^[SF]\d{2}$", re.IGNORECASE)  # S27, F27, S28 같은 시즌 표기


def extract_stage(sheet_name: str) -> str:
    """시트 탭 이름에서 샘플 단계(P1/P2/P3/GTM1/GTM2/SMS/PP)를 찾는다.
    여러 개가 동시에 있으면(예: 'GTM2-F27 PP 6.03') 이름에서 더 앞쪽에 나온 것을 사용한다.
    못 찾으면 시트 이름에서 시즌 표기(S27/F27 등)는 제외하고 그 다음 단어를 임시 라벨로 사용한다."""
    best = None
    for kw in STAGE_KEYWORDS:
        m = re.search(rf"\b{kw}\b", sheet_name, re.IGNORECASE)
        if m and (best is None or m.start() < best[1]):
            best = (kw, m.start())
    if best:
        return best[0]
    # 못 찾으면 시즌 표기(S27/F27)를 제외한 첫 영문 단어를 임시 라벨로 사용
    for m in re.finditer(r"[A-Za-z]{2,}\d{0,2}", sheet_name):
        token = m.group(0)
        if SEASON_TOKEN_RE.match(token):
            continue
        return token.upper()
    return "ETC"


@dataclass
class StyleInfo:
    season: str
    style_no: str
    description: str
    raw_filename: str


@dataclass
class DateBlock:
    """하나의 컬러웨이(colorway) x 하나의 제출일자 블록."""
    colorway: str
    date: Optional[dt.date]
    date_text: str
    fob: Optional[float]
    cm: Optional[float]
    cm_right: Optional[float]
    margin: Optional[float]
    remark: str
    remark_is_yellow: bool
    sheet_name: str
    header_row: int
    cm_row: int
    left_ttl_col: int
    right_ttl_col: int
    diffs: list = field(default_factory=list)  # 이전 날짜 대비 변경점 (문자열 리스트)
    sheet_order: int = 0  # 시트 탭 순서(값이 클수록 최신). 날짜가 같을 때 동점 처리용


def parse_filename(filename: str) -> StyleInfo:
    """파일명에서 시즌 / 7자리 스타일넘버 / DESCRIPTION 을 추출한다.

    예) '1789534333884_S27-1188891_M_FlyLux_Sweatpant-_fabTrim_updated_V-SUBMITTED.xlsx'
        -> season='S27', style_no='1188891', description='M FlyLux Sweatpant'
    """
    base = filename.rsplit("/", 1)[-1]
    base = re.sub(r"\.xlsx?$", "", base, flags=re.IGNORECASE)
    # 맨 앞 타임스탬프 숫자열 제거 (예: 1789534333884_)
    base = re.sub(r"^\d{6,}_", "", base)

    m = re.match(r"^([SF]\d{2}(?:[-/][SF]\d{2})?)[-_](\d{7})[_-]?(.*)$", base)
    if not m:
        # 혹시 순서가 다르면 느슨하게라도 시즌/스타일만 찾는다
        season_m = re.search(r"[SF]\d{2}(?:[-/][SF]\d{2})?", base)
        style_m = re.search(r"\d{7}", base)
        season = season_m.group(0) if season_m else ""
        style_no = style_m.group(0) if style_m else ""
        rest = base
    else:
        season, style_no, rest = m.groups()

    tokens = re.split(r"[_\-\s]+", rest)
    desc_tokens = []
    for tok in tokens:
        if not tok:
            continue
        if STOPWORDS.match(tok):
            break
        desc_tokens.append(tok)
    description = " ".join(desc_tokens).strip()

    return StyleInfo(season=season, style_no=style_no, description=description, raw_filename=filename)


def _cell_fill_rgb(cell) -> Optional[str]:
    try:
        fill = cell.fill
        if fill and fill.fgColor and fill.fgColor.type == "rgb":
            return fill.fgColor.rgb
    except Exception:
        pass
    return None


def _is_yellow(cell) -> bool:
    rgb = _cell_fill_rgb(cell)
    return bool(rgb) and rgb.upper() in YELLOW_FILLS


CBS_RE = re.compile(r"CBS\s*([0-9]{1,2}/[0-9]{1,2})\D*\$?\s*([0-9]+(?:\.[0-9]+)?)", re.IGNORECASE)


def _find_header_rows(ws):
    """'ITEM' 라벨이 있는 컬럼(D일)과 그 행(header_row)들을 모두 찾는다."""
    headers = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=min(ws.max_column, 10)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().upper() == "ITEM":
                headers.append((cell.row, cell.column))
    return headers


def _ttl_and_remark_cols(ws, header_row: int):
    ttl_cols, remark_cols = [], []
    for cell in ws[header_row]:
        if isinstance(cell.value, str):
            v = cell.value.strip().upper()
            if v == "TTL":
                ttl_cols.append(cell.column)
            elif v == "REMARK":
                remark_cols.append(cell.column)
    ttl_cols.sort()
    remark_cols.sort()
    return ttl_cols, remark_cols


def _find_margin(ws, header_row: int, cbs_row: Optional[int], left_ttl_col: int):
    """블록 최종 합계 행(=CBS가 적힌 행)에서 마진율(%)을 찾는다.
    마진율은 왼쪽(제안가) 블록의 'LOSS' 헤더 열(TTL 바로 왼쪽 열)에, CBS와 같은 행에
    '%' 서식으로 적혀 있다 (예: PP S27 7.14 시트의 H41, H84 등)."""
    if cbs_row is None:
        return None
    loss_col = left_ttl_col - 1
    header_val = ws.cell(row=header_row, column=loss_col).value
    if not (isinstance(header_val, str) and header_val.strip().upper() == "LOSS"):
        # 혹시 열 위치가 다르면 헤더 행에서 TTL 왼쪽에 있는 LOSS 열을 다시 탐색
        loss_col = None
        for c in range(1, left_ttl_col):
            v = ws.cell(row=header_row, column=c).value
            if isinstance(v, str) and v.strip().upper() == "LOSS":
                loss_col = c
        if loss_col is None:
            return None
    cell = ws.cell(row=cbs_row, column=loss_col)
    if not isinstance(cell.value, (int, float)):
        return None
    if "%" not in (cell.number_format or ""):
        return None
    return float(cell.value)


def _find_cm_row(ws, item_col: int, start_row: int, end_row: int):
    for r in range(start_row, end_row + 1):
        v = ws.cell(row=r, column=item_col).value
        if isinstance(v, str) and v.strip().upper() == "CM":
            return r
    return None


STYLE_NO_LEADING_RE = re.compile(r"^\s*(\d{7})\b")


def _sketch_color_col(ws, header_row: int, max_col: int = 8) -> Optional[int]:
    """헤더 행에서 'SKETCH/COLOR' 라벨이 있는 열을 찾는다."""
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str) and "SKETCH" in v.strip().upper():
            return c
    return None


def _declared_style_no(ws, header_row: int) -> Optional[str]:
    """블록의 SKETCH/COLOR 칸(예: '1184056 Trail Bottom ET PRT')맨 앞에 적힌
    7자리 스타일번호를 찾는다. 파일명의 스타일번호와 다르면 그 블록은 잘못 섞여
    들어간 것으로 보고 무시하기 위함.

    S/# 칸에는 'PRINT VERSION OF 1181433', 'REPEAT OF F26-1170233' 처럼 다른
    스타일번호를 참고용으로 언급하는 메모가 적히는 경우가 있어, 그 칸은 보지 않고
    SKETCH/COLOR 칸(숫자로 시작하는 셀)만 확인한다."""
    sc_col = _sketch_color_col(ws, header_row)
    if sc_col is None:
        return None
    for r in range(header_row + 1, header_row + 8):
        v = ws.cell(row=r, column=sc_col).value
        if v is None:
            continue
        m = STYLE_NO_LEADING_RE.match(str(v))
        if m:
            return m.group(1)
    return None


def _colorway_from_block(ws, header_row: int, end_row: int) -> str:
    """블록 내 원단 설명(주로 header_row 근처 R열 등)에서 컬러웨이를 추정한다.
    컬러 관련 표기가 전혀 없으면 SOLID로 본다 (PRINT/HEATHER는 항상 별도로
    명시적으로 표기되어 있다는 것이 확인된 규칙)."""
    texts = []
    for row in ws.iter_rows(min_row=max(1, header_row - 3), max_row=min(end_row, header_row + 6)):
        for cell in row:
            if isinstance(cell.value, str):
                texts.append(cell.value)
    blob = " ".join(texts).upper()
    if "HEATHER" in blob:
        return "HEATHER"
    if re.search(r"\bPRINT\b", blob) or re.search(r"\bPRT\b", blob):
        return "PRINT"
    if "SOLID" in blob:
        return "SOLID"
    return "SOLID" if header_row < end_row else "COLOR 2"


def parse_date_text(date_text: str, fallback_year: Optional[int] = None) -> Optional[dt.date]:
    m = re.match(r"^(\d{1,2})/(\d{1,2})$", date_text.strip())
    if not m:
        return None
    month, day = int(m.group(1)), int(m.group(2))
    year = fallback_year or dt.date.today().year
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def parse_sheet(ws, sheet_name: str, style_no: str):
    """워크시트 하나에서 컬러웨이별 DateBlock 리스트를 뽑는다.

    반환: (blocks, skip_notes) - skip_notes 는 파일명과 다른 스타일번호가 적혀 있어
    무시한 블록에 대한 안내 메시지 리스트."""
    blocks = []
    skip_notes = []
    header_rows = [r for r, c in _find_header_rows(ws)]
    if not header_rows:
        return blocks, skip_notes

    header_rows.sort()
    for idx, hr in enumerate(header_rows):
        item_col = 4  # 'ITEM' 은 통상 D열
        next_hr = header_rows[idx + 1] if idx + 1 < len(header_rows) else ws.max_row

        declared = _declared_style_no(ws, hr)
        if declared and style_no and declared != style_no:
            skip_notes.append(
                f"[{sheet_name}] 시트 내 스타일번호({declared})가 파일명 스타일번호({style_no})와 달라 해당 블록을 무시했습니다."
            )
            continue

        ttl_cols, remark_cols = _ttl_and_remark_cols(ws, hr)
        if not ttl_cols:
            continue
        left_ttl, right_ttl = ttl_cols[0], ttl_cols[-1]
        left_remark = remark_cols[0] if remark_cols else None
        right_remark = remark_cols[-1] if remark_cols else None

        cm_row = _find_cm_row(ws, item_col, hr + 1, next_hr - 1)
        if cm_row is None:
            continue

        cm_left = ws.cell(row=cm_row, column=left_ttl).value
        cm_right = ws.cell(row=cm_row, column=right_ttl).value if right_ttl != left_ttl else None

        remark_cell = ws.cell(row=cm_row, column=right_remark) if right_remark else None
        remark_text = remark_cell.value if remark_cell and isinstance(remark_cell.value, str) else ""
        remark_yellow = _is_yellow(remark_cell) if remark_cell is not None else False

        # CBS 날짜/가격 검색 (해당 블록 범위 내에서) - 찾은 행 번호도 함께 기억해둔다
        # (마진%이 CBS와 같은 행, 즉 그 블록의 최종 합계 행에 함께 적혀있기 때문)
        date_text, fob, cbs_row = "", None, None
        for row in ws.iter_rows(min_row=hr, max_row=next_hr - 1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, str):
                    m = CBS_RE.search(cell.value)
                    if m:
                        date_text = m.group(1)
                        cbs_row = cell.row
                        try:
                            fob = float(m.group(2))
                        except ValueError:
                            fob = None
                        break
            if date_text:
                break

        margin = _find_margin(ws, hr, cbs_row, left_ttl)

        colorway = _colorway_from_block(ws, hr, next_hr - 1)
        # 같은 시트에 같은 컬러웨이가 중복되면 구분자 추가
        existing = [b for b in blocks if b.colorway == colorway]
        if existing:
            colorway = f"{colorway} #{len(existing) + 1}"

        blocks.append(
            DateBlock(
                colorway=colorway,
                date=None,
                date_text=date_text,
                fob=fob,
                cm=cm_left if isinstance(cm_left, (int, float)) else None,
                cm_right=cm_right if isinstance(cm_right, (int, float)) else None,
                margin=margin,
                remark=remark_text.strip() if remark_text else "",
                remark_is_yellow=remark_yellow,
                sheet_name=sheet_name,
                header_row=hr,
                cm_row=cm_row,
                left_ttl_col=left_ttl,
                right_ttl_col=right_ttl,
            )
        )
    return blocks, skip_notes


def _blocks_structurally_compatible(ws_a, ws_b, header_row_a: int, header_row_b: int, sample_rows: int = 10) -> bool:
    """두 블록이 같은 템플릿(행 배열)인지 대략 확인. ITEM/Detail 라벨 열이 서로 다른
    버전(예: 옛 P1/P2/P3 라운드 시트)이면 행 정렬이 어긋나 오탐 diff가 나오므로 걸러낸다."""
    matches, total = 0, 0
    for offset in range(1, sample_rows + 1):
        va = ws_a.cell(row=header_row_a + offset, column=4).value
        vb = ws_b.cell(row=header_row_b + offset, column=4).value
        if va is None and vb is None:
            continue
        total += 1
        if isinstance(va, str) and isinstance(vb, str) and va.strip().lower() == vb.strip().lower():
            matches += 1
        elif va == vb:
            matches += 1
    if total == 0:
        return True
    return (matches / total) >= 0.6


def _header_name_map(ws, header_row: int, max_col: int = 20):
    """헤더 행에서 {(열이름, L/R): 열번호} 매핑을 만든다. CIF 열을 기준으로
    왼쪽(제안가) / 오른쪽(실제/개정가) 블록을 구분해서, 두 시트의 컬럼 배치가
    (K/L 열 유무 등으로) 달라도 '같은 이름의 같은 쪽 열'끼리 비교할 수 있게 한다."""
    names = {}
    cif_col = None
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str) and v.strip().upper() == "CIF" and cif_col is None:
            cif_col = c
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str) and v.strip():
            side = "R" if (cif_col and c >= cif_col) else "L"
            names[(v.strip().upper(), side)] = c
    return names


def _diff_block_cells(ws_a, ws_b, header_row_a: int, header_row_b: int, end_row_a: int):
    """같은 컬러웨이 블록에서 노란색 & 값이 달라진 셀만 비교. 컬럼 이름(헤더) 기준으로
    매칭하여, 두 시트의 열 배치가 달라도 정확히 대응되는 항목끼리 비교한다."""
    if not _blocks_structurally_compatible(ws_a, ws_b, header_row_a, header_row_b):
        return None  # 구조가 달라 비교 불가

    map_a = _header_name_map(ws_a, header_row_a)
    map_b = _header_name_map(ws_b, header_row_b)
    common_keys = set(map_a) & set(map_b)
    # REMARK 열은 별도 로직(remark 필드)에서 이미 다루므로 diff 노이즈에서 제외
    common_keys = {k for k in common_keys if k[0] not in ("REMARK",)}

    diffs = []
    max_row_offset = end_row_a - header_row_a
    for offset in range(1, max_row_offset + 1):
        r_a = header_row_a + offset
        r_b = header_row_b + offset
        label = None
        for c in (4, 5, 3):  # D=ITEM, E=Detail, C=FABRIC
            v = ws_a.cell(row=r_a, column=c).value
            if isinstance(v, str) and v.strip():
                label = v.strip()
                break
        for name, side in common_keys:
            c_a, c_b = map_a[(name, side)], map_b[(name, side)]
            cell_a = ws_a.cell(row=r_a, column=c_a)
            cell_b = ws_b.cell(row=r_b, column=c_b)
            va, vb = cell_a.value, cell_b.value
            if va == vb:
                continue
            if not (_is_yellow(cell_a) or _is_yellow(cell_b)):
                continue
            if va is None and vb is None:
                continue
            tag = "개정" if side == "R" else "제안"
            diffs.append((f"{label or f'row{r_a}'} ({name.title()}/{tag})", va, vb))
    return diffs


def parse_submitted_workbook(path: str, filename_for_parsing: Optional[str] = None):
    """SUBMITTED 워크북 전체를 파싱해서 컬러웨이별 시간순 DateBlock 리스트를 반환.

    반환: (style_info, {colorway: [DateBlock, ...(날짜 오름차순)]}, skip_notes)
    """
    fname = filename_for_parsing or path
    style_info = parse_filename(fname)

    wb = openpyxl.load_workbook(path, data_only=True)
    # 시트 탭 순서: 관찰된 실제 파일들은 "최신 -> 과거" 순으로 탭이 나열되어 있어
    # 역순으로 뒤집으면 과거->최신 시간 순서에 근접한다 (값이 클수록 최신/더 왼쪽 탭).
    # CBS 날짜가 비어있는 중간 협상 단계 시트가 있어 날짜만으로는 완전한 정렬이 어렵고,
    # 같은 날짜가 겹치는 경우의 동점 처리 기준으로도 사용한다.
    sheet_order = {name: idx for idx, name in enumerate(reversed(wb.sheetnames))}

    all_blocks_by_colorway: dict[str, list[DateBlock]] = {}
    all_skip_notes: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        blocks, skip_notes = parse_sheet(ws, sheet_name, style_info.style_no)
        all_skip_notes.extend(skip_notes)
        for b in blocks:
            b.sheet_order = sheet_order.get(sheet_name, 0)
            # normalize colorway name: strip numbering suffix used only for
            # within-sheet de-dup and merge back for cross-sheet grouping
            base_colorway = re.sub(r"\s*#\d+$", "", b.colorway)
            all_blocks_by_colorway.setdefault(base_colorway, []).append(b)

    # 날짜 파싱 및 정렬 (날짜가 있으면 날짜 우선, 없으면 시트 탭 순서로 보완;
    # 날짜가 같으면 시트 탭 순서(sheet_order)가 더 큰 쪽=더 왼쪽/최신 탭이 뒤로 온다)
    for colorway, blocks in all_blocks_by_colorway.items():
        for b in blocks:
            b.date = parse_date_text(b.date_text) if b.date_text else None
        blocks.sort(key=lambda b: (b.date or dt.date.min, b.sheet_order))

    # 단계별 차이점(diff) 계산 - 같은 컬러웨이 내 시간순으로 인접한 시트를 비교하고,
    # CBS 날짜가 없는 중간 단계의 변경점은 다음 확정(날짜 있는) 단계에 함께 모아 기재한다.
    for colorway, blocks in all_blocks_by_colorway.items():
        pending_notes: list[str] = []
        for i in range(1, len(blocks)):
            prev_b, cur_b = blocks[i - 1], blocks[i]
            if prev_b.sheet_name != cur_b.sheet_name:
                try:
                    ws_prev = wb[prev_b.sheet_name]
                    ws_cur = wb[cur_b.sheet_name]
                    end_row = prev_b.cm_row + 35  # 블록이 대략 CM행 기준 위아래 40행 내
                    diffs = _diff_block_cells(ws_prev, ws_cur, prev_b.header_row, cur_b.header_row, end_row)
                    if diffs is None:
                        pending_notes.append(f"[{prev_b.sheet_name}→{cur_b.sheet_name}] ⚠ 구조가 달라 자동 비교 생략 (직접 확인 필요)")
                    else:
                        for label, va, vb in diffs[:8]:
                            pending_notes.append(f"{label}: {va} -> {vb}")
                except Exception:
                    pass
            if cur_b.date and cur_b.fob is not None:
                cur_b.diffs = pending_notes
                pending_notes = []
        # 마지막까지 확정 안 된 변경점이 남아있으면 마지막 블록에 붙여준다
        if pending_notes and blocks:
            blocks[-1].diffs = (blocks[-1].diffs or []) + pending_notes

    return style_info, all_blocks_by_colorway, all_skip_notes
